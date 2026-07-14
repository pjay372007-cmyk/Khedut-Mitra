import os
import sys
import json
import csv
import re
import xml.etree.ElementTree as ET
import zipfile
from datetime import datetime

# Setup directories
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SOURCE_DIR = os.path.join(BASE_DIR, "knowledge_engine", "source_modules")
DATA_DIR = os.path.join(BASE_DIR, "data")
JS_DIR = os.path.join(os.path.dirname(BASE_DIR), "frontend", "scripts")
LOG_DIR = os.path.join(BASE_DIR, "logs")

os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(JS_DIR, exist_ok=True)
os.makedirs(LOG_DIR, exist_ok=True)

LOG_FILE = os.path.join(LOG_DIR, "knowledge_builder.log")

def log(msg, level="INFO"):
    timestamp = datetime.now().isoformat()
    log_line = f"[{timestamp}] [{level}] {msg}\n"
    try:
        print(log_line.strip())
    except UnicodeEncodeError:
        # Fallback for Windows CP1252 console printing
        try:
            encoding = sys.stdout.encoding or 'utf-8'
            print(log_line.strip().encode(encoding, errors='replace').decode(encoding))
        except Exception:
            # Fallback of fallback: print only ascii part
            print(log_line.strip().encode('ascii', errors='replace').decode('ascii'))
            
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(log_line)


# Common crop name translations for ID normalization
CROP_TRANSLATIONS = {
    "કપાસ": "cotton",
    "જુવાર": "jowar",
    "ચણા": "chana",
    "મગફળી": "groundnut",
    "ઘઉં": "wheat",
    "દિવેલ": "castor",
    "ડાંગર": "paddy",
    "શેરડી": "sugarcane",
    "બાજરી": "bajra",
    "તુવેર": "pigeonpea",
    "ચોળી": "cowpea",
    "રાયડો": "mustard"
}


# --- Extractors ---

def extract_pdf(file_path):
    log(f"Extracting PDF: {file_path}")
    text = ""
    try:
        import pdfplumber
        with pdfplumber.open(file_path) as pdf:
            for i, page in enumerate(pdf.pages):
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception as e:
        log(f"pdfplumber failed on {file_path}, trying PyPDF2: {e}", "WARNING")
        try:
            import PyPDF2
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
        except Exception as e2:
            log(f"Failed to extract PDF {file_path}: {e2}", "ERROR")
    return text


def extract_docx(file_path):
    log(f"Extracting DOCX: {file_path}")
    text = ""
    try:
        import docx
        doc = docx.Document(file_path)
        for para in doc.paragraphs:
            text += para.text + "\n"
    except Exception as e:
        log(f"Failed to extract DOCX {file_path}: {e}", "ERROR")
    return text


def extract_xlsx(file_path):
    log(f"Extracting Excel: {file_path}")
    rows = []
    # Try openpyxl
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active
        for r in sheet.iter_rows(values_only=True):
            rows.append([str(cell) if cell is not None else "" for cell in r])
        return rows
    except Exception as e:
        log(f"openpyxl not available or failed, using custom zipfile parser: {e}", "WARNING")

    # Zipfile XML fallback parser for XLSX
    try:
        shared_strings = []
        with zipfile.ZipFile(file_path) as z:
            if "xl/sharedStrings.xml" in z.namelist():
                with z.open("xl/sharedStrings.xml") as f:
                    tree = ET.parse(f)
                    root = tree.getroot()
                    ns = {'ns': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                    for t in root.findall('.//ns:t', ns):
                        shared_strings.append(t.text)

            with z.open("xl/worksheets/sheet1.xml") as f:
                tree = ET.parse(f)
                root = tree.getroot()
                ns = {'ns': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                for row_el in root.findall('.//ns:row', ns):
                    row_data = []
                    for cell in row_el.findall('ns:c', ns):
                        val_el = cell.find('ns:v', ns)
                        val = val_el.text if val_el is not None else ""
                        t_attr = cell.get('t')
                        if t_attr == 's' and val.isdigit():
                            idx = int(val)
                            if idx < len(shared_strings):
                                val = shared_strings[idx]
                        row_data.append(val)
                    rows.append(row_data)
    except Exception as e:
        log(f"Custom XLSX extraction failed for {file_path}: {e}", "ERROR")
    return rows


# --- Normalizer and Unstructured Document Ingest Engine ---

def parse_unstructured_text_to_records(text, category):
    """
    Parses unstructured text extracts (from PDFs/DOCX/TXT/MD) using regex heuristics
    to structure them into a normalized list of objects matching the category.
    """
    records = []
    
    # Prepend newline to ensure first heading matches split pattern
    text_to_split = "\n" + text
    blocks = re.split(r'(?i)\n\s*(?:crop|disease|pest|deficiency|scheme|faq)[:\-#\s]+', text_to_split)
    
    for block in blocks:
        if not block.strip():
            continue
            
        record = {}
        
        # Set default name/question as the first line of the block since the header tag was consumed
        lines = [line.strip() for line in block.split('\n') if line.strip()]
        if lines:
            record["name"] = lines[0]
            if category == "faqs":
                record["question"] = lines[0]
        
        # Specific fields matching
        name_match = re.search(r'(?i)name[:\- ]+([^\n]+)', block)
        if name_match:
            record["name"] = name_match.group(1).strip()
            
        name_gu_match = re.search(r'(?i)name_gu[:\- ]+([^\n]+)', block)
        if name_gu_match:
            record["name_gu"] = name_gu_match.group(1).strip()
            
        scientific_match = re.search(r'(?i)(?:scientific_name|scientific)[:\- ]+([^\n]+)', block)
        if scientific_match:
            record["scientific_name"] = scientific_match.group(1).strip()

        desc_match = re.search(r'(?i)(?:description|desc|info)[:\- ]+([^\n]+(?:\n\s+[^\n]+)*)', block)
        if desc_match:
            record["description"] = desc_match.group(1).strip()

        # Category specific parsing
        if category == "diseases":
            crop_match = re.search(r'(?i)crop[:\- ]+([^\n]+)', block)
            if crop_match:
                record["crop"] = [c.strip().lower() for c in crop_match.group(1).split(",")]
            severity_match = re.search(r'(?i)severity[:\- ]+(low|medium|moderate|severe|critical)', block)
            record["severity"] = severity_match.group(1).strip().capitalize() if severity_match else "Moderate"
            
            # Treatments parse
            treatments = []
            treat_matches = re.findall(r'(?i)treatment[:\- ]+([^\n]+)', block)
            for tm in treat_matches:
                treatments.append({"icon": "fa-flask-vial", "title": "Treatment Guide", "desc": tm.strip()})
            if treatments:
                record["treatment"] = treatments

        elif category == "faqs":
            q_match = re.search(r'(?i)(?:question|q)[:\- ]+([^\n]+)', block)
            a_match = re.search(r'(?i)(?:answer|a)[:\- ]+([^\n]+(?:\n\s+[^\n]+)*)', block)
            if q_match:
                record["question"] = q_match.group(1).strip()
            if a_match:
                record["answer"] = a_match.group(1).strip()
                
        # If we successfully extracted at least a name or question, add it!
        if record.get("name") or record.get("question"):
            records.append(record)
            
    return records


# --- Core Ingestion Pipeline ---

def build_category_database(category):
    log(f"Building Database for: {category}...")
    cat_dir = os.path.join(SOURCE_DIR, category)
    records = []
    
    if not os.path.exists(cat_dir):
        log(f"Directory {cat_dir} does not exist. Skipping.", "WARNING")
        return []
        
    for filename in os.listdir(cat_dir):
        file_path = os.path.join(cat_dir, filename)
        if filename == "info.txt":
            continue
            
        # Ingest JSON
        if filename.endswith(".json"):
            try:
                with open(file_path, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    if isinstance(data, list):
                        records.extend(data)
                    elif isinstance(data, dict):
                        # check if it wraps a list
                        for key in ["crops", "diseases", "pests", "deficiencies", "fertilizers", "organic", "schemes", "faqs", "advisories"]:
                            if key in data and isinstance(data[key], list):
                                records.extend(data[key])
                                break
                        else:
                            # Single dict
                            records.append(data)
                log(f"Successfully loaded JSON from {filename}")
            except Exception as e:
                log(f"Error loading JSON file {filename}: {e}", "ERROR")
                
        # Ingest CSV
        elif filename.endswith(".csv"):
            try:
                with open(file_path, "r", encoding="utf-8") as f:
                    reader = csv.DictReader(f)
                    for row in reader:
                        # Convert comma-separated columns if any
                        clean_row = {}
                        for k, v in row.items():
                            if k in ["crop", "region", "symptoms", "preventive"]:
                                clean_row[k] = [item.strip() for item in v.split(",") if item.strip()]
                            else:
                                clean_row[k] = v.strip() if v else ""
                        records.append(clean_row)
                log(f"Successfully loaded CSV from {filename}")
            except Exception as e:
                log(f"Error loading CSV file {filename}: {e}", "ERROR")

        # Ingest XLSX (Excel)
        elif filename.endswith(".xlsx"):
            rows = extract_xlsx(file_path)
            if len(rows) > 1:
                headers = [h.strip() for h in rows[0]]
                for row in rows[1:]:
                    record = {}
                    for idx, cell_val in enumerate(row):
                        if idx < len(headers):
                            record[headers[idx]] = cell_val.strip()
                    if record:
                        records.append(record)
                log(f"Successfully loaded Excel sheet from {filename}")

        # Ingest PDF
        elif filename.endswith(".pdf"):
            pdf_text = extract_pdf(file_path)
            extracted_records = parse_unstructured_text_to_records(pdf_text, category)
            records.extend(extracted_records)
            log(f"Extracted {len(extracted_records)} records from PDF: {filename}")

        # Ingest DOCX
        elif filename.endswith(".docx"):
            docx_text = extract_docx(file_path)
            extracted_records = parse_unstructured_text_to_records(docx_text, category)
            records.extend(extracted_records)
            log(f"Extracted {len(extracted_records)} records from DOCX: {filename}")

        # Ingest TXT / MD
        elif filename.endswith((".txt", ".md")):
            try:
                with open(file_path, "r", encoding="utf-8") as f:
                    text = f.read()
                extracted_records = parse_unstructured_text_to_records(text, category)
                records.extend(extracted_records)
                log(f"Extracted {len(extracted_records)} records from Text: {filename}")
            except Exception as e:
                log(f"Error reading Text file {filename}: {e}", "ERROR")

    # --- Deduplication, Normalization & Merging ---
    deduplicated = []
    seen_ids = set()

    for r in records:
        # 1. Map alternative ID keys
        if "disease_id" in r:
            r["id"] = r.pop("disease_id")
        elif "crop_id" in r:
            r["id"] = r.pop("crop_id")
        elif "pest_id" in r:
            r["id"] = r.pop("pest_id")
        elif "faq_id" in r:
            r["id"] = r.pop("faq_id")
            
        # 2. Extract and map alternative Name/Question keys
        r_name = ""
        if "name" in r:
            r_name = r["name"]
        elif "Crop_Name" in r:
            r_name = r["Crop_Name"]
            r["name"] = r_name
        elif "Pest_Name" in r:
            r_name = r["Pest_Name"]
            r["name"] = r_name
        elif "disease_name" in r:
            if isinstance(r["disease_name"], dict):
                r_name = r["disease_name"].get("english", "")
            else:
                r_name = r["disease_name"]
            r["name"] = r_name
        elif "crop_name_english" in r:
            r_name = r["crop_name_english"]
            r["name"] = r_name
            
        # Handle FAQ structures (questions / answers nested)
        if "questions" in r and isinstance(r["questions"], dict):
            r["question"] = r["questions"].get("english", "")
            r["question_gu"] = r["questions"].get("gujarati", "")
            r_name = r["question"]
        elif "question" in r:
            if isinstance(r["question"], dict):
                r["question_gu"] = r["question"].get("gujarati", "")
                r["question"] = r["question"].get("english", "")
            r_name = r["question"]
            
        if "answer" in r and isinstance(r["answer"], dict):
            r["answer_gu"] = r["answer"].get("gujarati", "")
            r["answer"] = r["answer"].get("english", "")
            
        # Handle disease-specific nested fields
        if category == "diseases":
            # Map namegu
            if "disease_name" in r and isinstance(r["disease_name"], dict):
                r["namegu"] = r["disease_name"].get("gujarati", r.get("namegu", ""))
            # Normalize crop list
            if "crop" in r:
                if isinstance(r["crop"], str):
                    r["crop"] = [c.strip().lower() for c in r["crop"].split(",") if c.strip()]
            else:
                r["crop"] = []

        # 3. Resolve IDs with smart translation fallback
        r_id = r.get("id")
        if not r_id and r_name:
            # Try to map crop name Gujarati characters to safe English keys
            mapped_name = ""
            for gu_word, en_word in CROP_TRANSLATIONS.items():
                if gu_word in r_name:
                    mapped_name = en_word
                    break
            
            if mapped_name:
                r_id = f"{category}_{mapped_name}"
            else:
                clean_name = re.sub(r'[^a-z0-9]+', '_', r_name.lower().strip())
                clean_name = re.sub(r'_+', '_', clean_name).strip('_')
                if clean_name:
                    r_id = f"{category}_{clean_name}"
                else:
                    # Fallback to index count
                    r_id = f"{category}_{len(seen_ids) + 1}"
            
            r["id"] = r_id
            
        if not r_id:
            continue
            
        r_id = r_id.strip()
        r["id"] = r_id
        
        # Merge logic
        if r_id in seen_ids:
            existing = next(item for item in deduplicated if item["id"] == r_id)
            log(f"Merging duplicate record for ID: {r_id}")
            for key, val in r.items():
                if not existing.get(key):
                    existing[key] = val
                elif isinstance(existing[key], list) and isinstance(val, list):
                    existing[key] = list(set(existing[key] + val))
                elif isinstance(existing[key], dict) and isinstance(val, dict):
                    existing[key].update(val)
        else:
            seen_ids.add(r_id)
            deduplicated.append(r)

    log(f"Database build complete. Unique records count: {len(deduplicated)}")
    return deduplicated


# --- Validation Engine ---

def validate_database_records(records, category):
    validation_report = {
        "category": category,
        "total_records": len(records),
        "errors": [],
        "warnings": [],
        "passed": 0
    }
    
    # Required keys mapping per category
    required_keys = {
        "crops": ["id", "name"],
        "diseases": ["id", "name", "description"],
        "pests": ["id", "name"],
        "faqs": ["id", "question", "answer"]
    }
    
    reqs = required_keys.get(category, ["id"])
    
    for r in records:
        r_id = r.get("id", "UnknownID")
        r_name = r.get("name") or r.get("question") or "Unnamed"
        
        # Check required fields
        missing_fields = [f for f in reqs if f not in r or not r[f]]
        if missing_fields:
            msg = f"Record '{r_name}' ({r_id}) is missing fields: {missing_fields}"
            validation_report["errors"].append({"id": r_id, "message": msg})
            log(msg, "ERROR")
            continue
            
        # Check translations (Bilingual verification)
        has_gujarati = False
        if category == "crops":
            has_gujarati = bool(r.get("crop_name_gujarati") or r.get("name_gu") or r.get("crop_name_gujarati_unicode") or r.get("Crop_Name"))
        elif category == "diseases":
            has_gujarati = bool(r.get("namegu") or r.get("disease_name", {}).get("gujarati"))
        elif category == "faqs":
            has_gujarati = bool(r.get("question_gu") or r.get("question_gujarati"))
        else:
            has_gujarati = any(k.endswith("_gu") or k.endswith("gujarati") for k in r.keys()) or bool(r.get("name_gu"))

        if not has_gujarati:
            msg = f"Record '{r_name}' ({r_id}) does not have a Gujarati translation."
            validation_report["warnings"].append({"id": r_id, "message": msg})
            log(msg, "WARNING")

        validation_report["passed"] += 1

    return validation_report


# --- Main Coordinator ---

def main():
    log("=========================================")
    log("Starting KrishiAI Production Knowledge Builder Pipeline")
    log("=========================================")
    
    categories = [
        "crops", "diseases", "pests", "deficiencies", "fertilizers",
        "organic", "schemes", "weather", "market", "faqs", "general"
    ]
    
    compiled_kb = {
        "metadata": {
            "version": "3.0.0",
            "release_date": datetime.now().strftime("%Y-%m-%d"),
            "author": "KrishiAI Knowledge Engine Team",
            "region": "Gujarat, India"
        }
    }
    
    reports = []
    
    for cat in categories:
        records = build_category_database(cat)
        compiled_kb[cat] = records
        
        # Save modular json files
        output_file = os.path.join(DATA_DIR, f"{cat}.json")
        with open(output_file, "w", encoding="utf-8") as f:
            json.dump(records, f, indent=2, ensure_ascii=False)
        log(f"Saved modular database to: {output_file}")
        
        # Validate
        rep = validate_database_records(records, cat)
        reports.append(rep)

    # In addition, to support backward-compatibility with index.html loading KB.diseases
    # we need to map compiled diseases to match exactly the format of standard KB.diseases!
    # Let's align keys:
    aligned_diseases = []
    for d in compiled_kb["diseases"]:
        aligned_d = d.copy()
        
        # Map crop names inside diseases to match standard CROP_CLASSES array
        if "crop" in d and isinstance(d["crop"], list):
            cleaned_crops = []
            for crop_str in d["crop"]:
                crop_lower = crop_str.lower()
                if "cotton" in crop_lower:
                    cleaned_crops.append("cotton")
                elif "wheat" in crop_lower or "ઘઉં" in crop_lower:
                    cleaned_crops.append("wheat")
                elif "groundnut" in crop_lower or "મગફળી" in crop_lower:
                    cleaned_crops.append("groundnut")
                elif "rice" in crop_lower or "paddy" in crop_lower or "ડાંગર" in crop_lower:
                    cleaned_crops.append("paddy")
                elif "sugarcane" in crop_lower or "શેરડી" in crop_lower:
                    cleaned_crops.append("sugarcane")
                elif "castor" in crop_lower or "દિવેલ" in crop_lower:
                    cleaned_crops.append("castor")
                else:
                    cleaned_crops.append(crop_lower)
            aligned_d["crop"] = cleaned_crops

        # Support default fields if they are missing
        if not aligned_d.get("type"):
            aligned_d["type"] = "fungal"
        if not aligned_d.get("urgency"):
            aligned_d["urgency"] = "medium"
        if not aligned_d.get("emoji"):
            aligned_d["emoji"] = "🔬"
            
        aligned_diseases.append(aligned_d)
    
    # Save the consolidated JSON file
    consolidated_path = os.path.join(DATA_DIR, "knowledge_base.json")
    with open(consolidated_path, "w", encoding="utf-8") as f:
        json.dump(compiled_kb, f, indent=2, ensure_ascii=False)
    log(f"Saved consolidated JSON database to: {consolidated_path}")

    # Generate js/knowledge_base.js for client-side inclusion (declares global window.KB)
    js_kb_file = os.path.join(JS_DIR, "knowledge_base.js")
    js_content = f"""/**
 * KrishiAI Unified Knowledge Base Auto-Generated Database
 * Generated on: {datetime.now().isoformat()}
 */
window.KB = {{
  diseases: {json.dumps(aligned_diseases, indent=2, ensure_ascii=False)},
  crops: {json.dumps(compiled_kb["crops"], indent=2, ensure_ascii=False)},
  pests: {json.dumps(compiled_kb["pests"], indent=2, ensure_ascii=False)},
  faqs: {json.dumps(compiled_kb["faqs"], indent=2, ensure_ascii=False)},
  cropTips: {{
    wheat: [
      "CRI irrigation is most critical. Never skip it.",
      "Apply urea in 2 splits — at sowing and at CRI stage."
    ],
    cotton: [
      "Avoid pesticides during flowering. They kill honeybees.",
      "Monitor for whitefly from 30 DAS."
    ],
    groundnut: [
      "Never disturb soil during pegging stage (30–45 DAS).",
      "Gypsum application at pegging is essential for pod development."
    ],
    paddy: [
      "Maintain 2–5 cm standing water in first 30 days.",
      "Apply zinc sulphate if Khaira disease symptoms appear."
    ],
    cumin: [
      "Blight is main threat in cumin — spray Mancozeb preventively.",
      "Cumin needs cool dry weather — avoid late sowing."
    ],
    sugarcane: [
      "Ratoon management — remove dry leaves monthly.",
      "Red rot is the biggest disease — use disease-free setts."
    ],
    default: [
      "Soil testing every 3 years is the best investment.",
      "Crop rotation reduces pest and disease pressure significantly."
    ]
  }}
}};
"""
    with open(js_kb_file, "w", encoding="utf-8") as f:
        f.write(js_content)
    log(f"Saved client-side JS knowledge base to: {js_kb_file}")

    # Generate a unified validation report file
    val_report_path = os.path.join(DATA_DIR, "validation_report.json")
    with open(val_report_path, "w", encoding="utf-8") as f:
        json.dump(reports, f, indent=2, ensure_ascii=False)
    log(f"Saved validation report to: {val_report_path}")

    log("=========================================")
    log("Knowledge Builder Pipeline Execution Successful!")
    log("=========================================")


if __name__ == "__main__":
    main()
